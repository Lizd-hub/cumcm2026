"""Compile report source and verified results into a Chinese DOCX."""
import json
import os
import re
from pathlib import Path
import fitz
from PIL import Image
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from native_math import equations as native_equations, insert_equation

BASE = Path(__file__).resolve().parent
TMP = BASE.parent / 'tmp' / 'b_report_assets'
OUT = BASE.parent / 'outputs'
TMP.mkdir(parents=True, exist_ok=True)
OUT.mkdir(exist_ok=True)
WIDTH = 6.94


def md_table(headers, rows):
    return '\n'.join(['| ' + ' | '.join(headers) + ' |', '| ' + ' | '.join(['---'] * len(headers)) + ' |']
                     + ['| ' + ' | '.join(map(str, r)) + ' |' for r in rows])


def compile_source():
    summary = json.loads((BASE / 'results' / 'summary.json').read_text())
    def find(q, model, scenario='random_smooth'):
        return next(r for r in summary if r['question'] == q and r['model'] == model and r['scenario'] == scenario)
    replacements = {}
    for q in (3, 4):
        b, n = find(q, 'baseline'), find(q, 'improved')
        replacements[f'Q{q}_BASE'] = f"{b['mean_case_average_s']:.2f}"
        replacements[f'Q{q}_NEW'] = f"{n['mean_case_average_s']:.2f}"
        replacements[f'Q{q}_GAIN'] = f"{100*(1-n['mean_case_average_s']/b['mean_case_average_s']):.2f}"
    audits = json.loads((BASE / 'results' / 'data_audit.json').read_text())
    replacements['AUDIT_TABLE'] = md_table(['示例', '原始动作数', '有效示向度', '按协议缺省', '异常记录', '填补个数'],
        [[f"第{r['question']}问 " + ('常规' if r['scenario']=='random_smooth' else '压力'), r['raw_records'],
          r['direction'], r['structurally_absent_bearings'], r['invalid_records'], 0] for r in audits])
    candidates = json.loads((BASE / 'results' / 'q2_candidates.json').read_text())
    selected = [(750,500,'基础示例'), (750,650,'本次候选中半径最小'), (600,800,'候选区边界上的例子')]
    q2rows = []
    for x, y, desc in selected:
        r = next(c for c in candidates if c['x']==x and c['y']==y)
        q2rows.append([f'({x}, {y})', f"{r['move_s']:.2f}", f"{r['sample_worst_radius_m']:.2f}", desc])
    replacements['Q2_TABLE'] = md_table(['第二点局部坐标 m', '移动时间 s', '样本最大包围半径 m', '解释'], q2rows)
    replacements['RESULTS_TABLE'] = md_table(['问题', '场景', '策略', '局数', '清除比例', '每源均值 s', '最差局均值 s'],
        [[r['question'], '常规' if r['scenario']=='random_smooth' else '联合压力',
          '基础' if r['model']=='baseline' else '改进', r['cases'], f"{r['clear_fraction']:.0%}",
          f"{r['mean_case_average_s']:.2f}", f"{r['max_case_average_s']:.2f}"] for r in summary])
    replacements['COST_TABLE'] = md_table(['常规场景', '平均路程 km', '平均检测次数', '平均失败清除次数'],
        [[f"第{r['question']}问 " + ('基础' if r['model']=='baseline' else '改进'),
          f"{r['mean_distance_m']/1000:.2f}", f"{r['mean_measurements']:.2f}", f"{r['mean_failed_clears']:.2f}"]
         for r in summary if r['scenario']=='random_smooth'])
    source = (BASE / 'report_template.md').read_text(encoding='utf-8')
    for key, value in replacements.items():
        source = source.replace('{{' + key + '}}', value)
    if re.search(r'\{\{[^}]+\}\}', source):
        raise ValueError('Unresolved report template field')
    (BASE / 'report.md').write_text(source, encoding='utf-8')
    return source


def east_asia_font(element, east='宋体', latin='Times New Roman'):
    rpr = element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts')
        rpr.insert(0, fonts)
    for key in ('ascii', 'hAnsi', 'cs'):
        fonts.set(qn('w:' + key), latin)
    fonts.set(qn('w:eastAsia'), east)
    for key in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        fonts.attrib.pop(qn('w:' + key), None)


def set_run(run, size=None, bold=None, mono=False):
    east_asia_font(run._element, latin='DejaVu Sans Mono' if mono else 'Times New Roman')
    run.font.color.rgb = RGBColor(0,0,0)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def text_runs(paragraph, text, size=None):
    for piece in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text):
        if not piece:
            continue
        mono = piece.startswith('`') and piece.endswith('`')
        bold = piece.startswith('**') and piece.endswith('**')
        if mono:
            piece = piece[1:-1]
        if bold:
            piece = piece[2:-2]
        run = paragraph.add_run(piece)
        set_run(run, size=size, bold=bold if bold else None, mono=mono)
        if paragraph.style.name in ('Title','Heading 1','Heading 2','Heading 3'):
            east_asia_font(run._element, east='黑体')


def table_widths(headers):
    n = len(headers)
    if headers[0] == '题意或规则':
        return [2.3, .95, WIDTH-3.25]
    if n == 7:
        return [.43,.86,.55,.44,.79,1.33,WIDTH-4.4]
    if n == 6:
        return [1.29,1.13,1.13,1.13,1.13,WIDTH-5.81]
    if n == 5:
        return [.50,1.05,1.38,1.80,WIDTH-4.73]
    if n == 4 and headers[0] == '小问':
        return [.65,2.02,2.05,WIDTH-4.72]
    if n == 4 and headers[0] == '检测点':
        return [1.1,1.95,1.95,WIDTH-5.0]
    if n == 4 and headers[0].startswith('第二点'):
        return [1.7,1.08,1.9,WIDTH-4.68]
    if n == 3 and headers[0] == '检测点':
        return [1.25,(WIDTH-1.25)/2,(WIDTH-1.25)/2]
    if n == 3:
        return [1.45,2.35,WIDTH-3.8]
    return [WIDTH/n]*n


def add_table(doc, lines):
    rows = [[s.strip() for s in line.strip().strip('|').split('|')] for line in lines]
    rows = [r for r in rows if not all(re.fullmatch(r':?-+:?', v) for v in r)]
    widths = table_widths(rows[0])
    table = doc.add_table(rows=0, cols=len(rows[0]))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    props = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        child = OxmlElement('w:' + edge)
        for key,val in (('val','single'),('sz','5'),('color','D9D9D9')):
            child.set(qn('w:'+key),val)
        borders.append(child)
    props.append(borders)
    margins = OxmlElement('w:tblCellMar')
    for edge, value in (('top','85'),('bottom','85'),('left','95'),('right','95')):
        child = OxmlElement('w:' + edge)
        child.set(qn('w:w'),value)
        child.set(qn('w:type'),'dxa')
        margins.append(child)
    props.append(margins)
    for index, values in enumerate(rows):
        row = table.add_row()
        trpr = row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if index == 0:
            trpr.append(OxmlElement('w:tblHeader'))
        for cell,width,value in zip(row.cells,widths,values):
            cell.width = Inches(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.first_line_indent = Inches(0)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.08
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if index==0 else WD_ALIGN_PARAGRAPH.LEFT
            text_runs(p,value,size=9.5)
            if index==0:
                for run in p.runs:
                    run.bold = True
                shade = OxmlElement('w:shd')
                shade.set(qn('w:fill'),'EAF2F8')
                cell._tc.get_or_add_tcPr().append(shade)
    tail = doc.add_paragraph()
    tail.paragraph_format.space_after = Pt(1)
    tail.paragraph_format.space_before = Pt(0)
    tail.paragraph_format.line_spacing = 0.3
    set_run(tail.add_run(''),size=3)


def main():
    source = compile_source()
    # Bundle-provided Droid CJK is used only for local rendering. Word keeps
    # standard Chinese font names and falls back on the reader's system.
    font_path = TMP / 'DroidSansFallback.ttf'
    if not font_path.exists():
        font_path.write_bytes(fitz.Font('cjk').buffer)
    config = f'''<?xml version="1.0"?><!DOCTYPE fontconfig SYSTEM "fonts.dtd"><fontconfig>
<include ignore_missing="yes">/etc/fonts/fonts.conf</include><dir>{TMP}</dir><cachedir>{TMP / 'fontcache'}</cachedir>
<dir>/opt/codex/runtimes/codex-primary-runtime/dependencies/native/libreoffice-headless/libreoffice/share/fonts/truetype</dir>
<alias><family>Cambria Math</family><prefer><family>DejaVu Math TeX Gyre</family></prefer></alias>
<alias><family>宋体</family><prefer><family>Droid Sans Fallback</family></prefer></alias>
<alias><family>SimSun</family><prefer><family>Droid Sans Fallback</family></prefer></alias>
<alias><family>黑体</family><prefer><family>Droid Sans Fallback</family></prefer></alias>
</fontconfig>'''
    (TMP / 'fontconfig.xml').write_text(config, encoding='utf-8')
    equations = re.findall(r'\$\$\s*\n(.*?)\n\$\$', source, flags=re.S)
    expressions = native_equations()
    assert len(expressions) == len(equations), 'Native equations must match report source order.'
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(.78)
    sec.top_margin = Inches(.70)
    sec.bottom_margin = Inches(.68)
    sec.header_distance, sec.footer_distance = Inches(.28), Inches(.28)
    normal = doc.styles['Normal']
    east_asia_font(normal._element)
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor(0,0,0)
    normal.paragraph_format.line_spacing = 1.16
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.first_line_indent = Pt(22)
    for name,size in [('Title',21),('Subtitle',12),('Heading 1',15),('Heading 2',12.5),('Heading 3',11.5)]:
        style = doc.styles[name]
        east_asia_font(style._element, east='黑体' if name!='Subtitle' else '宋体')
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = name!='Subtitle'
        style.font.color.rgb = RGBColor(0,0,0)
        style.paragraph_format.first_line_indent = Pt(0)
        style.paragraph_format.space_before = Pt(12 if name.startswith('Heading') else 5)
        style.paragraph_format.space_after = Pt(7)
        style.paragraph_format.keep_with_next = True
    header = sec.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.first_line_indent = Pt(0)
    text_runs(header,'2026 B题 数学建模详解',size=9)
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Pt(0)
    text_runs(footer,'第 ',size=9)
    run = footer.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'),'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'),'preserve'); instr.text=' PAGE '
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'),'end')
    run._r.extend([begin,instr,end]); set_run(run,size=9)
    text_runs(footer,' 页',size=9)
    lines = source.splitlines()
    i, eqindex = 0, 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith('# '):
            p = doc.add_paragraph(style='Title')
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            text_runs(p,line[2:])
        elif line.startswith('## '):
            p = doc.add_paragraph(style='Heading 1'); text_runs(p,line[3:])
        elif line.startswith('### '):
            p = doc.add_paragraph(style='Heading 2'); text_runs(p,line[4:])
        elif i < 6 and line.startswith('2026'):
            p=doc.add_paragraph(style='Subtitle'); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; text_runs(p,line)
        elif line.startswith('|'):
            group = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                group.append(lines[i]); i += 1
            add_table(doc,group); continue
        elif line == '$$':
            i += 1
            while i < len(lines) and lines[i].strip() != '$$':
                i += 1
            p = doc.add_paragraph()
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(8)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            insert_equation(p, expressions[eqindex])
            eqindex += 1
        elif line.startswith('!['):
            match = re.fullmatch(r'!\[(.*)\]\((.*)\)',line)
            imgpath = BASE/match.group(2)
            w,h=Image.open(imgpath).size
            width=min(6.4,4.25*w/h)
            p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.first_line_indent=Pt(0)
            p.paragraph_format.keep_with_next=True
            p.paragraph_format.space_after=Pt(4)
            pic=p.add_run().add_picture(str(imgpath),width=Inches(width))
            pic._inline.docPr.set('descr',match.group(1))
            caption=doc.add_paragraph(); caption.alignment=WD_ALIGN_PARAGRAPH.CENTER
            caption.paragraph_format.first_line_indent=Pt(0)
            caption.paragraph_format.space_after=Pt(9)
            text_runs(caption,match.group(1),size=9.5)
        elif line.startswith('```'):
            i += 1
            while i < len(lines) and not lines[i].strip().startswith('```'):
                p=doc.add_paragraph(); p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.space_after=Pt(2); p.paragraph_format.line_spacing=1.05
                set_run(p.add_run(lines[i]),size=8.5,mono=True)
                i += 1
        else:
            p=doc.add_paragraph()
            text_runs(p,line)
            if line.startswith('[') and re.match(r'\[\d\]',line):
                p.paragraph_format.first_line_indent=Pt(0)
                p.paragraph_format.line_spacing=Pt(12)
                p.paragraph_format.space_after=Pt(2)
                for run in p.runs:
                    run.font.size=Pt(9.5)
        i += 1
    assert eqindex == len(equations)
    # No theme-driven heading colors or paragraph rules remain.
    for style in doc.styles:
        if style.name in ('Title','Subtitle','Heading 1','Heading 2','Heading 3'):
            for color in style._element.iter(qn('w:color')):
                color.attrib.clear(); color.set(qn('w:val'),'000000')
            for border in list(style._element.iter(qn('w:pBdr'))):
                border.getparent().remove(border)
    doc.core_properties.title='无线电干扰源自动定位与清除建模详解'
    doc.core_properties.subject='2026 B题完整解题方案与逐问讲解'
    doc.core_properties.author=''
    target=OUT/'B题_完整解答与逐问讲解.docx'
    doc.save(target)
    print(json.dumps({'document':str(target),'equations':eqindex,'source_characters':len(source),
                      'fontconfig':str(TMP/'fontconfig.xml')},ensure_ascii=False))


if __name__ == '__main__':
    main()
