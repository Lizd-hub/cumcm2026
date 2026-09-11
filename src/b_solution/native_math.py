"""Explicit native Word equations for this report, with structured fractions,
subscripts, superscripts and radicals. No rasterization or external math library.
The resulting OMML is inspected through the DOCX renderer before delivery.
"""
from copy import deepcopy
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def run(text):
    r = OxmlElement('m:r')
    pr = OxmlElement('m:rPr')
    sty = OxmlElement('m:sty'); sty.set(qn('m:val'), 'p')
    pr.append(sty); r.append(pr)
    wr = OxmlElement('w:rPr')
    fonts = OxmlElement('w:rFonts')
    fonts.set(qn('w:ascii'), 'Cambria Math'); fonts.set(qn('w:hAnsi'), 'Cambria Math')
    wr.append(fonts)
    size = OxmlElement('w:sz'); size.set(qn('w:val'), '23'); wr.append(size)
    r.append(wr)
    t = OxmlElement('m:t'); t.set(qn('xml:space'), 'preserve'); t.text = text
    r.append(t)
    return r


def nodes(value):
    if isinstance(value, str):
        return [run(value)]
    if isinstance(value, (list, tuple)):
        out=[]
        for v in value:
            out.extend(nodes(v))
        return out
    return [deepcopy(value)]


def slot(tag, value):
    e=OxmlElement(tag)
    e.extend(nodes(value))
    return e


def sub(base, lower):
    e=OxmlElement('m:sSub'); e.append(slot('m:e',base)); e.append(slot('m:sub',lower)); return e


def sup(base, upper):
    e=OxmlElement('m:sSup'); e.append(slot('m:e',base)); e.append(slot('m:sup',upper)); return e


def frac(num, den):
    e=OxmlElement('m:f'); e.append(slot('m:num',num)); e.append(slot('m:den',den)); return e


def root(value):
    e=OxmlElement('m:rad'); p=OxmlElement('m:radPr'); h=OxmlElement('m:degHide')
    h.set(qn('m:val'),'1'); p.append(h); e.append(p); e.append(OxmlElement('m:deg')); e.append(slot('m:e',value)); return e


def lower_op(name, lower):
    e=OxmlElement('m:limLow'); e.append(slot('m:e',name)); e.append(slot('m:lim',lower)); return e


def array(rows):
    e=OxmlElement('m:eqArr')
    for row in rows:
        e.append(slot('m:e',row))
    return e


def equations():
    return [
        ['Ω = {G ∈ ',sup('ℝ','2'), ': ‖G‖ ≤ 1800},     1000 ≤ R ≤ 1500.'],
        ['u(φ) = (cos φ, sin φ),     [a,b] = ',sub('a','x'),sub('b','y'),' − ',sub('a','y'),sub('b','x'),'.'],
        array([
            ['[u(',sub('θ','i'),' − δ), G − ',sub('S','i'),'] ≥ 0,'],
            ['[u(',sub('θ','i'),' + δ), G − ',sub('S','i'),'] ≤ 0.']]),
        ['D(P) = ',lower_op('max','G,H ∈ P'),' ‖G − H‖ = ',lower_op('max','1 ≤ j,k ≤ m'),
         ' ‖',sub('V','j'),' − ',sub('V','k'),'‖.'],
        array([[sub('r','*'),' = ',lower_op('min','c'),lower_op('max','G ∈ P'),' ‖G − c‖,'],
               [sub('r','*'),' ≤ 20  ⇒  ‖G − ',sub('c','*'),'‖ ≤ 20   (∀G ∈ P).']]),
        array([
            ['𝒞 = B(S,1000) ∩ B(S + 1000u(θ − δ),1000)'],
            ['∩ B(S + 1000u(θ + δ),1000).']]),
        [sup('‖Q − S − du‖','2'),' − ',sup('d','2'),' = ',sup('‖Q − S‖','2'),' − 2d(Q − S) · u ≤ 0.'],
        [sub('Q','*'),' ∈ ',lower_op('arg min',['Q ∈ ',sub('𝒞','cand')]),' {',
         frac('‖Q − S‖','5'),' + 5 + λ',lower_op('max','ω ∈ 𝒲'),' r(',sub('P','2'),'(Q,ω))}.'],
        ['T = ',frac('L','5'),' + 5',sub('N','m'),' + ',sub('N','s'),' + 5',sub('N','c'),' + 3',sub('N','f'),'.'],
        [sub('S','0'),' = (0,0),     ',sub('S','k+1'),' = 900',root('3'),' u(60°k),     k = 0,1,…,5.'],
        array([[sup('d','2'),' ≤ ',sup('ρ','2'),' + ',sup('a','2'),' − 2aρ cos 30°'],
               ['= ',sup('ρ','2'),' + 2430000 − 2700ρ.']]),
        [sub('d','max'),' = ',frac(['25',root('2')],'2'),' ≈ 17.68 m < 20 m.'],
        ['‖S − G‖ ≤ R,     u(α) · (S − G) ≥ 0.'],
        array([[sub('S','ij'),' = a (i + ',frac('j','2'),', ',frac(root('3'),'2'),' j),'],
               ['i,j ∈ ℤ,     ‖',sub('S','ij'),'‖ ≤ 1800 + a.']]),
        array([['G = ',sub('λ','1'),sub('S','1'),' + ',sub('λ','2'),sub('S','2'),' + ',sub('λ','3'),sub('S','3'),','],
               [sub('λ','i'),' ≥ 0,     ',sub('λ','1'),' + ',sub('λ','2'),' + ',sub('λ','3'),' = 1.']])
    ]


def insert_equation(paragraph, expression):
    math=OxmlElement('m:oMath')
    math.extend(nodes(expression))
    paragraph._p.append(math)


if __name__=='__main__':
    from pathlib import Path
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.shared import Inches
    target=Path(__file__).resolve().parent.parent/'tmp'/'b_report_assets'/'native_math_check.docx'
    doc=Document()
    section=doc.sections[0]
    section.page_width=Inches(8.5); section.page_height=Inches(11)
    section.left_margin=section.right_margin=Inches(.75)
    for i,expr in enumerate(equations()):
        doc.add_paragraph(f'Equation {i+1}')
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        insert_equation(p,expr)
    doc.save(target)
    print(target)
